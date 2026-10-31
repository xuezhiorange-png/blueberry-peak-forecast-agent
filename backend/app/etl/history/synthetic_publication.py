"""Clean-room public examples. No private reader, redaction or prediction adapter."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import random
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from backend.app.etl.history.fingerprint import business_fingerprint
from backend.app.etl.history.normalizer import normalize_text
from backend.app.etl.history.parser import parse_date_value, parse_weight_value
from backend.app.etl.history.schemas import EXPECTED_HEADERS, FatalQualityThresholds, ImportRules

DATA_CLASS = "SYNTHETIC_ONLY"
SOURCE = "TD01C_INDEPENDENT_GENERATOR_V1"
PRIVATE_REAL_POLICY = {
    "public_export": False,
    "reader_implemented": False,
    "preserve": "original dates, kg, identities, relationships, revisions and availability",
    "future_training": "requires separate source, PIT, label and execution admission",
}
FIELD_POLICIES = {
    key: {"private": "PRESERVE_UNCHANGED_PRIVATE", "public": treatment}
    for key, treatment in zip(
        EXPECTED_HEADERS,
        (
            "FICTIONAL_DATE",
            "FICTIONAL_LINK",
            "FICTIONAL_FARM",
            "FICTIONAL_SUBFARM",
            "FICTIONAL_VARIETY",
            "SYNTHETIC_GRADE",
            "INDEPENDENT_GENERATED_KG",
            "FICTIONAL_FACTORY",
        ),
        strict=True,
    )
}
FIELD_POLICIES.update(
    {
        key: {"private": "RETAIN_PRIVATE", "public": "OMIT_OR_GENERATE_FRESH"}
        for key in (
            "filename",
            "sheet",
            "business_id",
            "coordinates",
            "area",
            "supplier_relation",
            "derived_totals",
            "daily_curve",
            "author",
            "comments",
            "links",
        )
    }
)
_RULES = ImportRules(
    version=SOURCE,
    valid_months=set(range(1, 13)),
    excluded_grades=set(),
    excluded_factories=set(),
    deduplicate_suspected_business_rows_in_curated=False,
    date_formats=["%Y-%m-%d"],
    variety_prefixes_to_remove=[],
    empty_strings={""},
    max_issue_examples=0,
    allow_unknown_factory_in_analysis=False,
    allow_unknown_variety_in_analysis=False,
    allow_empty_factory_in_analysis=False,
    allow_empty_variety_in_analysis=False,
    fatal_quality_thresholds=FatalQualityThresholds(),
)
_ZIP_MEMBERS = {
    "docProps/app.xml",
    "docProps/core.xml",
    "xl/theme/theme1.xml",
    "xl/worksheets/sheet1.xml",
    "xl/styles.xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "[Content_Types].xml",
}
_STAMP = datetime(2000, 1, 1)


@dataclass(frozen=True)
class PublicRow:
    season: str
    values: tuple[str, ...]
    issues: tuple[str, ...]


@dataclass(frozen=True)
class PublicPackage:
    seed: int
    rows: tuple[PublicRow, ...]
    data_class: str = DATA_CLASS
    source_identity: str = SOURCE


def _generate(seed: int) -> PublicPackage:
    if type(seed) is not int or not 0 <= seed < 2**31:
        raise ValueError("BOUNDED_INTEGER_SEED_REQUIRED")
    rng = random.Random(seed)
    rows = []
    for year in (2040, 2041):
        season = f"{year - 1}-{year}"
        values_list = []
        for day in range(12):
            for farm in (1, 2):
                for variety in (1, 2):
                    weight = format(Decimal(rng.randrange(1, 900000)) / Decimal(1000), ".6f")
                    values_list.append(
                        [
                            (date(year, 1, 1) + timedelta(days=day)).isoformat(),
                            f"SYNTHETIC_LINK_{farm}",
                            f"SYNTHETIC_FARM_{farm}",
                            f"SYNTHETIC_FARM_{farm}_SUB_1",
                            f"SYNTHETIC_VARIETY_{variety}",
                            "SYNTHETIC_GRADE_A",
                            weight,
                            f"SYNTHETIC_FACTORY_{farm}",
                        ]
                    )
        values_list[0][6] = ""
        values_list[1][6] = "0.000000"
        values_list[2][6] = "-1.250000"
        values_list[3][0] = "SYNTHETIC_UNKNOWN_DATE"
        values_list[4][5] = "SYNTHETIC_GRADE_EXCLUDED"
        values_list.append(list(values_list[5]))
        seen = set()
        for values in values_list:
            parsed_date, date_errors = parse_date_value(values[0], None, 0, _RULES)
            weight_kg, weight_errors = parse_weight_value(values[6], _RULES)
            issues = [*date_errors, *weight_errors]
            if weight_kg is not None:
                if not weight_kg.is_finite():
                    raise ValueError("NONFINITE_GENERATED_WEIGHT")
                if weight_kg == 0:
                    issues.append("zero_weight")
                if weight_kg < 0:
                    issues.append("negative_weight")
            if values[5] == "SYNTHETIC_GRADE_EXCLUDED":
                issues.append("synthetic_excluded_grade")
            fingerprint = business_fingerprint(
                season_code=season,
                receipt_date=parsed_date,
                factory_name=values[7],
                farm_name=values[2],
                subfarm_name=values[3],
                variety_name=values[4],
                grade_code=values[5],
                weight_kg=weight_kg,
            )
            if fingerprint in seen:
                issues.append("suspected_duplicate")
            seen.add(fingerprint)
            rows.append(PublicRow(season, tuple(values), tuple(sorted(issues))))
    return PublicPackage(seed, tuple(rows))


def generate(
    seed: int = 41, *, source_path: Path | None = None, source_policy: str = "GENERATED_ONLY"
) -> PublicPackage:
    # Refuse before exists/open/hash. A name, label or local file is not provenance.
    if source_path is not None or source_policy != "GENERATED_ONLY":
        raise ValueError("REAL_SOURCE_READ_NOT_AUTHORIZED")
    return _generate(seed)


def validate_public(package: PublicPackage) -> None:
    if package != _generate(package.seed):
        raise ValueError("NOT_GENERATOR_OUTPUT")
    for row in package.rows:
        if len(row.values) != len(EXPECTED_HEADERS):
            raise ValueError("EXPORT_SCHEMA_MISMATCH")
        if any(normalize_text(value) != (value or None) for value in row.values):
            raise ValueError("EXPORT_NORMALIZATION_MISMATCH")


def require_usage(package: PublicPackage, usage: str) -> None:
    validate_public(package)
    if usage not in {"TEST", "DEMO"}:
        raise ValueError("SYNTHETIC_NOT_A_FORMAL_SOURCE")


def canonical_payload(package: PublicPackage) -> bytes:
    validate_public(package)
    payload = {
        "data_class": package.data_class,
        "source_identity": package.source_identity,
        "seed": package.seed,
        "headers": EXPECTED_HEADERS,
        "quantity_basis": "SYNTHETIC_DECLARED_LEGACY_RECEIPT_KG",
        "rows": [
            {"season": r.season, "values": r.values, "issues": r.issues} for r in package.rows
        ],
        "source_equivalence_proven": False,
        "strict_pit_validated": False,
        "available_at": None,
        "formal_source_admitted": False,
    }
    return (json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def _season_rows(package: PublicPackage, season: str) -> tuple[PublicRow, ...]:
    validate_public(package)
    rows = tuple(row for row in package.rows if row.season == season)
    if not rows:
        raise ValueError("UNKNOWN_SYNTHETIC_SEASON")
    return rows


def xlsx_bytes(package: PublicPackage, season: str = "2039-2040") -> bytes:
    rows = _season_rows(package, season)
    book = Workbook()
    sheet = book.active
    sheet.title = DATA_CLASS
    book.properties.creator = DATA_CLASS
    book.properties.lastModifiedBy = DATA_CLASS
    book.properties.created = _STAMP
    book.properties.modified = _STAMP
    sheet.append(EXPECTED_HEADERS)
    for row in rows:
        sheet.append(row.values)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    # openpyxl updates modified at save. Canonicalise only this new workbook.
    out = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(out, "w") as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content, count = re.subn(
                    rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    content,
                )
                if count != 2:
                    raise ValueError("MISSING_SYNTHETIC_TIMESTAMP")
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, content)
    data = out.getvalue()
    check_xlsx(data, package, season)
    return data


def check_csv(data: bytes, package: PublicPackage, season: str) -> None:
    expected = _season_rows(package, season)
    actual = list(csv.reader(io.StringIO(data.decode("utf-8"), newline="")))
    if actual != [list(EXPECTED_HEADERS), *(list(row.values) for row in expected)]:
        raise ValueError("UNSAFE_PUBLIC_CSV_CONTENT_BINDING")


def check_xlsx(data: bytes, package: PublicPackage, season: str = "2039-2040") -> None:
    expected_rows = _season_rows(package, season)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if (
            set(archive.namelist()) != _ZIP_MEMBERS
            or len(archive.infolist()) != 9
            or sum(info.file_size for info in archive.infolist()) > 2_000_000
            or archive.comment
        ):
            raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
        for name in archive.namelist():
            text = archive.read(name)
            if b"<!DOCTYPE" in text or b"<!ENTITY" in text:
                raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
            xml = ElementTree.fromstring(text)
            if any(node.get("TargetMode") == "External" for node in xml.iter()):
                raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
    book = load_workbook(io.BytesIO(data), keep_links=False)
    try:
        if book.sheetnames != [DATA_CLASS] or book.defined_names:
            raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
        for field in book.properties.__elements__:
            expected = {
                "creator": DATA_CLASS,
                "lastModifiedBy": DATA_CLASS,
                "created": _STAMP,
                "modified": _STAMP,
            }.get(field)
            if getattr(book.properties, field) != expected:
                raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
        sheet = book.active
        if sheet.sheet_state != "visible":
            raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
        actual = [
            tuple(cell.value if cell.value is not None else "" for cell in row) for row in sheet
        ]
        if actual != [EXPECTED_HEADERS, *(row.values for row in expected_rows)]:
            raise ValueError("UNSAFE_PUBLIC_WORKBOOK_CONTENT_BINDING")
        for row in sheet:
            for cell in row:
                if cell.data_type == "f" or cell.comment or cell.hyperlink:
                    raise ValueError("UNSAFE_PUBLIC_WORKBOOK")
    finally:
        book.close()


def export(package: PublicPackage, destination: Path) -> dict[str, str]:
    require_usage(package, "DEMO")
    outputs: dict[str, bytes] = {}
    for season in sorted({row.season for row in package.rows}):
        rows = tuple(row for row in package.rows if row.season == season)
        name = "SYNTHETIC_ONLY_" + season.replace("-", "_")
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(EXPECTED_HEADERS)
        writer.writerows(row.values for row in rows)
        outputs[name + ".csv"] = stream.getvalue().encode()
        check_csv(outputs[name + ".csv"], package, season)
        outputs[name + ".xlsx"] = xlsx_bytes(package, season)
    hashes = {name: hashlib.sha256(data).hexdigest() for name, data in outputs.items()}
    manifest = json.loads(canonical_payload(package))
    manifest["exports_sha256"] = hashes
    outputs["SYNTHETIC_ONLY-manifest.json"] = (
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    ).encode()
    destination.mkdir(mode=0o700, parents=True, exist_ok=False)
    for name, data in outputs.items():
        with os.fdopen(
            os.open(destination / name, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), "wb"
        ) as f:
            f.write(data)
    return hashes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("operation", choices=["generate"])
    parser.add_argument("--seed", type=int, default=41)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(export(generate(args.seed), args.output_dir), sort_keys=True))


if __name__ == "__main__":
    main()
