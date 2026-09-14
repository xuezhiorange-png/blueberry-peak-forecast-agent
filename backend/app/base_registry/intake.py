"""Read-only workbook intake and optional representative DEM enrichment."""

import json
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import urlopen

import openpyxl

from backend.app.area_yield.data import digest
from backend.app.area_yield.experiment import file_hash
from backend.app.base_registry.registry import area_value, build_registry


def import_workbook(path: Path, expected_hash: str, expected_count: int = 39) -> dict[str, Any]:
    if file_hash(path) != expected_hash:
        raise ValueError("workbook hash mismatch")
    original = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(original.worksheets) != 1:
            raise ValueError("unexpected sheets")
        values = list(original.active.iter_rows(values_only=True))
        cache = list(cached.active.iter_rows(values_only=True))
        if values[0] != ("基地名称", "覆盖农场", "经纬度", "亩数"):
            raise ValueError("workbook schema mismatch")
        rows = [list(r) for r in values[1:] if any(v is not None for v in r)]
        if len(rows) != expected_count:
            raise ValueError("base count mismatch")
        for row, cache_row in zip(values[1:], cache[1:], strict=True):
            if isinstance(row[3], str) and row[3].startswith("="):
                if area_value(row[3]) is None or area_value(row[3]) != area_value(cache_row[3]):
                    raise ValueError("unsupported area formula or cache mismatch")
        return build_registry(rows, expected_hash)
    finally:
        original.close()
        cached.close()


def fetch_elevation(registry: dict[str, Any]) -> dict[str, Any]:
    bases = [b for b in registry["bases"] if b["longitude"] is not None]
    query = {
        "longitude": ",".join(b["longitude"] for b in bases),
        "latitude": ",".join(b["latitude"] for b in bases),
    }
    url = "https://api.open-meteo.com/v1/elevation?" + urlencode(query)
    result: dict[str, Any] = {
        "provider": "Open-Meteo / Copernicus DEM GLO-90 2021",
        "reference": "https://open-meteo.com/en/docs/elevation-api",
        "method": "GET /v1/elevation supplied lon/lat without coordinate transformation",
        "coordinate_crs_caveat": "SOURCE_CRS_UNCONFIRMED_WGS84_QUERY_ASSUMPTION",
        "retrieved_at": datetime.now(UTC).isoformat(),
        "url": url,
        "rows": [],
    }
    try:
        with urlopen(url, timeout=40) as response:
            raw = response.read().decode()
        values = json.loads(raw)["elevation"]
        if len(values) != len(bases):
            raise ValueError("elevation count mismatch")
        result["raw_response"] = raw
        for base, value in zip(bases, values, strict=True):
            number = Decimal(str(value)) if value is not None else None
            valid = number is not None and number.is_finite()
            result["rows"].append(
                {
                    "base_id": base["base_id"],
                    "longitude": base["longitude"],
                    "latitude": base["latitude"],
                    "raw_value": value,
                    "elevation_m": str(number) if valid else None,
                    "status": "DEM_ESTABLISHED_CRS_REVIEW_REQUIRED" if valid else "NOT_ESTABLISHED",
                }
            )
    except (OSError, ValueError, KeyError, TypeError, ArithmeticError) as exc:
        result["failure_type"] = type(exc).__name__
        result["rows"] = [
            {
                "base_id": b["base_id"],
                "longitude": b["longitude"],
                "latitude": b["latitude"],
                "raw_value": None,
                "elevation_m": None,
                "status": "NOT_ESTABLISHED",
            }
            for b in bases
        ]
    result["hash"] = digest(result)
    return result


def point_in_ring(x: float, y: float, ring: list[list[float]]) -> bool:
    inside = False
    for a, b in zip(ring, ring[1:] + ring[:1], strict=True):
        if (a[1] > y) != (b[1] > y) and x < (b[0] - a[0]) * (y - a[1]) / (b[1] - a[1]) + a[0]:
            inside = not inside
    return inside


def point_in_geometry(x: float, y: float, geometry: dict[str, Any]) -> bool:
    if geometry["type"] not in {"Polygon", "MultiPolygon"}:
        raise ValueError("polygon required")
    polygons = (
        [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    )
    return any(
        point_in_ring(x, y, rings[0]) and not any(point_in_ring(x, y, r) for r in rings[1:])
        for rings in polygons
    )


def fetch_region_screen(registry: dict[str, Any]) -> dict[str, Any]:
    """Public province polygon for sanity screening only, not climate authority."""
    url = "https://geo.datav.aliyun.com/areas_v3/bound/530000.json"
    result: dict[str, Any] = {
        "url": url,
        "retrieved_at": datetime.now(UTC).isoformat(),
        "purpose": "REGION_SANITY_SCREEN_NOT_LEGAL_BOUNDARY_OR_CRS_VERIFICATION",
    }
    try:
        with urlopen(url, timeout=40) as response:
            raw = response.read().decode()
        geometry = json.loads(raw)["features"][0]["geometry"]
        result["raw_response"] = raw
        result["rows"] = [
            {
                "base_id": b["base_id"],
                "region_scope": (
                    "YUNNAN_CORE"
                    if point_in_geometry(float(b["longitude"]), float(b["latitude"]), geometry)
                    else "OUT_OF_YUNNAN"
                )
                if b["longitude"] is not None
                else "LOCATION_REVIEW_REQUIRED",
            }
            for b in registry["bases"]
        ]
    except (OSError, ValueError, KeyError, TypeError, IndexError) as exc:
        result["failure_type"] = type(exc).__name__
        result["rows"] = [
            {"base_id": b["base_id"], "region_scope": b["region_scope"]} for b in registry["bases"]
        ]
    result["hash"] = digest(result)
    return result
