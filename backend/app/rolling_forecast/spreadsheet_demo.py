"""Explicit test-environment workbook/source binding, never a user-file importer.

Only server-owned SYNTHETIC resources are read. Requests carry digests and column
indices, never uploaded cells, paths or quantities. Existing capabilities govern
all reads and issuance. Disabled in every non-test environment.
"""

import base64
import csv
import hashlib
import io
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import Field

from backend.app.local_auth.sessions import AuthError, Principal
from backend.app.rolling_forecast import application, inbound
from backend.app.rolling_forecast.policy import Catalog, Context, InboundSnapshot
from backend.app.rolling_forecast.schemas import Contract, Hash, Issue
from backend.app.rolling_forecast.source_readiness import inspect_source
from backend.app.rolling_forecast.source_readiness_report import REASONS
from backend.research.v019_synthetic_validation import digest

FIELDS = (
    "date",
    "base",
    "farm",
    "plot",
    "variety",
    "weight",
    "unit",
    "grade",
    "record_id",
    "recorded_at",
    "received_at",
    "revision_id",
)
SCENARIOS = Literal["complete", "missing-day", "late-record", "basis-conflict"]
ORIGINS = ("2026-11-05", "2026-11-06")


class Asset(Contract):
    path: str
    sha256: Hash
    sheet: str
    header: int = Field(ge=0, le=10)
    mapping: dict[str, int]


class Demo(Contract):
    context: str
    context_hash: Hash
    source_reference: str
    model_reference: str
    source_hash: Hash
    fit_hash: Hash
    projection_hash: Hash
    binding_hash: Hash
    assets: dict[Literal["xlsx", "csv"], Asset]


class Manifest(Contract):
    schema_version: Literal["PRE_DATA_SYNTHETIC_DEMO_V1"]
    data_class: Literal["SYNTHETIC_TEST_FIXTURE"]
    scenarios: dict[SCENARIOS, Demo]


class Binding(Contract):
    scenario: SCENARIOS
    format: Literal["csv", "xlsx"]
    binding_hash: Hash
    file_sha256: Hash
    projection_hash: Hash
    sheet: str = Field(max_length=100)
    header: int = Field(ge=0, le=10000)
    mapping: dict[str, int] = Field(max_length=12)
    declared_unit: Literal["kg"]
    declared_basis: Literal["RECORDED_INBOUND_WEIGHT"]
    origin: Literal["2026-11-05", "2026-11-06"] = "2026-11-05"


def projection_bytes(rows: list[list[str]]) -> bytes:
    return json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode()


def source_rows(context: Context, source: InboundSnapshot) -> list[list[str]]:
    """Fixed demo observations only, not the later forecast target labels."""
    index = inbound.read_index(context, source)
    rows = []
    for item in sorted(index.days, key=lambda d: d.business_date):
        if item.business_date > date(2026, 11, 5):
            continue
        day = inbound.Day.model_validate_json(
            inbound.read(Path(source.manifest_path).parent / item.file, item.sha256, 1024 * 1024)
        )
        if (day.base_id, day.season, day.business_date, day.source_system) != (
            context.area.base_id,
            context.calendar.season,
            item.business_date,
            source.source_system,
        ):
            raise ValueError("DEMO_SOURCE_IDENTITY_MISMATCH")
        for r in day.records:
            rows.append(
                [
                    str(day.business_date),
                    day.base_id,
                    r.member.farm_business_key,
                    r.member.subfarm_business_key,
                    r.member.variety_business_key,
                    format(r.quantity_kg, ".6f"),
                    "kg",
                    r.grade,
                    r.logical_id,
                    r.source_recorded_at.isoformat() if r.source_recorded_at else "",
                    r.server_ingested_at.isoformat() if r.server_ingested_at else "",
                    r.revision_id,
                ]
            )
    return rows


def workbook_rows(raw: bytes, kind: str, asset: Asset) -> list[list[str]]:
    if kind == "csv":
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    else:
        # Only a small, hash-pinned server-created workbook reaches this reader.
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book[asset.sheet]
            if sheet.max_row > 100 or sheet.max_column > 12:
                raise ValueError("DEMO_WORKBOOK_BOUNDS")
            rows = list(sheet.values)
        finally:
            book.close()
    result = []
    for row in rows[asset.header + 1 :]:
        converted = []
        for field in FIELDS:
            value = row[asset.mapping[field]]
            if field == "date":
                value = value.date().isoformat() if isinstance(value, datetime) else str(value)
            elif field == "weight":
                value = format(Decimal(str(value)), ".6f")
            converted.append("" if value is None else str(value).strip())
        result.append(converted)
    return result


def resolve(
    settings: Any, policy: Catalog, principal: Principal, scenario: SCENARIOS, operation: str
) -> tuple[Demo, Context, InboundSnapshot]:
    if (
        settings.app_env != "test"
        or not settings.pre_data_demo_manifest_path
        or not settings.pre_data_demo_manifest_sha256
    ):
        raise AuthError("SYNTHETIC_DEMO_UNAVAILABLE", 503)
    # Authorization before any source/workbook/model content. The manifest is quantity-free.
    allowed = application.contexts(policy, principal, operation)
    manifest = Manifest.model_validate_json(
        inbound.read(
            Path(settings.pre_data_demo_manifest_path),
            settings.pre_data_demo_manifest_sha256,
            65536,
        )
    )
    demo = manifest.scenarios.get(scenario)
    if demo is None:
        raise ValueError("INVALID_DEMO_SCENARIO")
    context = next((c for c in allowed if c.reference == demo.context), None)
    if context is None:
        raise AuthError("FORBIDDEN")
    if digest(context.model_dump(mode="json")) != demo.context_hash:
        raise ValueError("DEMO_CONTEXT_MISMATCH")
    source = next((s for s in context.snapshots if s.reference == demo.source_reference), None)
    model = next((m for m in context.models if m.reference == demo.model_reference), None)
    if (
        not isinstance(source, InboundSnapshot)
        or model is None
        or source.provenance != "SYNTHETIC_TEST_FIXTURE"
        or model.provenance != "SYNTHETIC_TEST_FIXTURE"
        or not source.source_system.startswith("SYNTHETIC")
    ):
        raise ValueError("REAL_SOURCE_EXECUTION_NOT_AUTHORIZED")
    if (source.snapshot_hash, model.fit_sha256) != (demo.source_hash, demo.fit_hash):
        raise ValueError("DEMO_SOURCE_IDENTITY_MISMATCH")
    return demo, context, source


def asset_bytes(demo: Demo, kind: str) -> bytes:
    asset = demo.assets[kind]  # type: ignore[index]
    return inbound.read(Path(asset.path), asset.sha256, 2 * 1024 * 1024)


def download(demo: Demo, kind: str) -> dict[str, Any]:
    asset = demo.assets[kind]  # type: ignore[index]
    return dict(
        data_class="SYNTHETIC_TEST_FIXTURE",
        binding_hash=demo.binding_hash,
        file_sha256=asset.sha256,
        projection_hash=demo.projection_hash,
        source_hash=demo.source_hash,
        sheet=asset.sheet,
        header=asset.header,
        mapping=asset.mapping,
        filename="SYNTHETIC-inbound." + kind,
        base64=base64.b64encode(asset_bytes(demo, kind)).decode(),
    )


def validate_binding(
    demo: Demo, context: Context, source: InboundSnapshot, request: Binding
) -> dict[str, Any]:
    asset = demo.assets[request.format]
    if (
        request.binding_hash,
        request.file_sha256,
        request.projection_hash,
        request.sheet,
        request.header,
        request.mapping,
    ) != (
        demo.binding_hash,
        asset.sha256,
        demo.projection_hash,
        asset.sheet,
        asset.header,
        asset.mapping,
    ):
        raise ValueError("UNREGISTERED_PREVIEW_FILE")
    rows = source_rows(context, source)
    if (
        hashlib.sha256(projection_bytes(rows)).hexdigest() != demo.projection_hash
        or workbook_rows(asset_bytes(demo, request.format), request.format, asset) != rows
    ):
        raise ValueError("DEMO_WORKBOOK_SOURCE_MISMATCH")
    origin = date.fromisoformat(request.origin)
    report = inspect_source(
        context,
        source.reference,
        demo.model_reference,
        origin,
        datetime.fromisoformat(request.origin + "T17:00:00+08:00"),
        mode="synthetic-full",
    )
    return dict(
        data_class="SYNTHETIC_TEST_FIXTURE",
        binding_hash=demo.binding_hash,
        readiness=report,
        reason_messages=[
            REASONS.get(code, "请核对来源合同与完整性证明") for code in report["blockers"]
        ],
    )


def issue_request(context: Context, demo: Demo, request: Binding) -> Issue:
    return Issue(
        base_id=context.area.base_id,
        target_season=context.area.season,
        origin_date=date.fromisoformat(request.origin),
        origin_cutoff_at=datetime.fromisoformat(request.origin + "T17:00:00+08:00").astimezone(UTC),
        observation_snapshot_reference=demo.source_reference,
        area_authority_reference=context.area.reference,
        frozen_model_artifact_reference=demo.model_reference,
        execution_mode="RESEARCH_SHADOW",
    )
