"""Deterministic public synthetic workbooks bound to actual S9/S8 sources.

No default capabilities; only an isolated test fixture may explicitly register
its synthetic principal. This is not a business import or auto-enrollment tool.
"""

import csv
import hashlib
import io
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook

from backend.app.rolling_forecast import spreadsheet_demo as demo
from backend.app.rolling_forecast.policy import Catalog
from backend.app.rolling_forecast.source_readiness_report import new_directory
from backend.research.v019_synthetic_validation import digest
from backend.tests.rolling_forecast.inbound_fixture import write
from backend.tests.rolling_forecast.source_readiness_package import create_package

HEADERS = (
    "到货业务日",
    "研究基地",
    "农场代号",
    "地块代号",
    "蓝莓品种",
    "登记重量",
    "计量单位",
    "商品等级",
    "业务流水号",
    "来源落笔时间",
    "服务接收时间",
    "修订版本",
)


def workbook(rows: list[list[str]]) -> bytes:
    book = Workbook()
    book.active.title = "说明"
    book.active.append(["SYNTHETIC_TEST_FIXTURE — 合成示例，不代表实际产量"])
    sheet = book.create_sheet("合成入库明细")
    sheet.append(["合成入库记录重量 · 非田间采收量"])
    sheet.append(["前三步仅本地预览，固定来源通过后才允许合成演示"])
    sheet.append(HEADERS)
    for row in rows:
        values = list(row)
        values[0] = datetime.fromisoformat(values[0])
        values[5] = int(values[5].split(".")[0])
        sheet.append(values)
    for row in sheet.iter_rows(min_row=4):
        row[0].number_format = "yyyy-mm-dd"
        row[5].number_format = "0.000000"
    book.properties.creator = "SYNTHETIC_TEST_FIXTURE"
    book.properties.created = book.properties.modified = datetime(2000, 1, 1)
    temp = io.BytesIO()
    book.save(temp)
    output = io.BytesIO()
    # Stable archive timestamps and core metadata, independent of build time/path.
    with ZipFile(temp) as src, ZipFile(output, "w") as dst:
        for name in sorted(src.namelist()):
            raw = src.read(name)
            if name == "docProps/core.xml":
                import re

                raw = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2000-01-01T00:00:00Z", raw)
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            dst.writestr(info, raw)
    return output.getvalue()


def create_demo(root: Path, principal_id: str | None = None) -> tuple[Catalog, Path, Path]:
    root = new_directory(root)
    contexts, capabilities, scenarios = [], [], {}
    for scenario in ("complete", "missing-day", "late-record", "basis-conflict"):
        catalog, _ = create_package(root / scenario, scenario, principal_id=principal_id)
        context = catalog.contexts[0]
        source = context.snapshots[0]
        rows = demo.source_rows(context, source)
        csv_file = io.StringIO(newline="")
        writer = csv.writer(csv_file)
        writer.writerow(HEADERS)
        writer.writerows(rows)
        assets = {}
        for kind, raw in (
            ("xlsx", workbook(rows)),
            ("csv", ("\ufeff" + csv_file.getvalue()).encode()),
        ):
            path = root / scenario / ("SYNTHETIC-inbound." + kind)
            path.write_bytes(raw)
            path.chmod(0o600)
            assets[kind] = dict(
                path=str(path.resolve()),
                sha256=hashlib.sha256(raw).hexdigest(),
                sheet="合成入库明细" if kind == "xlsx" else "CSV",
                header=2 if kind == "xlsx" else 0,
                mapping=dict(zip(demo.FIELDS, range(12), strict=True)),
            )
        projection_hash = hashlib.sha256(demo.projection_bytes(rows)).hexdigest()
        binding = digest(
            dict(
                source=source.snapshot_hash,
                fit=context.models[0].fit_sha256,
                files={k: v["sha256"] for k, v in assets.items()},
                projection=projection_hash,
            )
        )
        value = context.model_dump(mode="json")
        value["reference"] = "SYNTHETIC-PRE-DATA-" + scenario
        value["snapshots"][0]["reference"] = "SYNTHETIC-WORKBOOK-" + binding
        contexts.append(value)
        for cap in catalog.capabilities:
            capabilities.append({**cap.model_dump(mode="json"), "contexts": [value["reference"]]})
        scenarios[scenario] = dict(
            context=value["reference"],
            context_hash=digest(value),
            source_reference=value["snapshots"][0]["reference"],
            model_reference=context.models[0].reference,
            source_hash=source.snapshot_hash,
            fit_hash=context.models[0].fit_sha256,
            projection_hash=projection_hash,
            binding_hash=binding,
            assets=assets,
        )
    combined = Catalog.model_validate(
        dict(
            schema_version="ROLLING_SHADOW_CATALOG_V1",
            revision="SYNTHETIC-PRE-DATA-R1",
            contexts=contexts,
            capabilities=capabilities,
        )
    )
    path = root / "SYNTHETIC-catalog.json"
    write(path, combined.model_dump(mode="json"))
    manifest = demo.Manifest.model_validate(
        dict(
            schema_version="PRE_DATA_SYNTHETIC_DEMO_V1",
            data_class="SYNTHETIC_TEST_FIXTURE",
            scenarios=scenarios,
        )
    )
    manifest_path = root / "SYNTHETIC-workbooks.json"
    write(manifest_path, manifest.model_dump(mode="json"))
    return combined, path, manifest_path
