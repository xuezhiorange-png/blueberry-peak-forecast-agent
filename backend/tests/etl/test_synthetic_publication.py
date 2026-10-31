"""Independent artificial records, never renamed or copied private sources."""

import csv
import hashlib
import io
import json
import subprocess
import sys
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment

from backend.app.etl.history import synthetic_publication as publication
from backend.app.etl.history.schemas import EXPECTED_HEADERS

pytestmark = [pytest.mark.unit, pytest.mark.contract]


def test_independent_generation_relationships_quality_and_schema():
    package = publication.generate(41)
    assert package == publication.generate(41)
    assert package != publication.generate(42)
    assert len({row.season for row in package.rows}) == 2
    assert len({row.values[2] for row in package.rows}) == 2
    assert all(row.values[3].startswith(row.values[2] + "_") for row in package.rows)
    assert {code for row in package.rows for code in row.issues} >= {
        "empty_weight",
        "zero_weight",
        "negative_weight",
        "suspected_duplicate",
    }
    publication.validate_public(package)
    assert publication.FIELD_POLICIES.keys() >= set(EXPECTED_HEADERS)
    assert publication.PRIVATE_REAL_POLICY["public_export"] is False


@pytest.mark.parametrize("usage", ["TRAINING", "FORMAL_FORECAST", "REAL_SCORING", "AUTHORITY"])
def test_synthetic_is_not_a_real_source_or_production_permission(usage):
    with pytest.raises(ValueError, match="SYNTHETIC_NOT_A_FORMAL_SOURCE"):
        publication.require_usage(publication.generate(), usage)


@pytest.mark.parametrize("policy", ["PRIVATE_REAL", "SYNTHETIC_ONLY", "RENAMED_PRIVATE"])
def test_arbitrary_input_files_are_rejected_before_io(tmp_path, monkeypatch, policy):
    def forbidden(*args, **kwargs):
        raise AssertionError("input path must not be inspected or opened")

    monkeypatch.setattr(Path, "open", forbidden)
    monkeypatch.setattr(Path, "exists", forbidden)
    with pytest.raises(ValueError, match="REAL_SOURCE_READ_NOT_AUTHORIZED"):
        publication.generate(41, source_path=tmp_path / "renamed.xlsx", source_policy=policy)


@pytest.mark.parametrize(
    "field,value",
    [
        (2, "PRIVATE_NAME"),
        (0, "2025-01-01"),
        (6, "987654.123456"),
        (1, '=HYPERLINK("https://invalid")'),
    ],
)
def test_sensitive_fields_and_copied_curve_cannot_be_relabelled(field, value):
    package = publication.generate()
    row = package.rows[0]
    values = list(row.values)
    values[field] = value
    bad = replace(package, rows=(replace(row, values=tuple(values)), *package.rows[1:]))
    with pytest.raises(ValueError, match="NOT_GENERATOR_OUTPUT"):
        publication.validate_public(bad)


def test_private_source_identity_cannot_be_labelled_synthetic():
    package = publication.generate()
    with pytest.raises(ValueError, match="NOT_GENERATOR_OUTPUT"):
        publication.validate_public(replace(package, source_identity="SOURCE_002"))


def test_export_and_independent_process_bytes_are_identical(tmp_path):
    first = tmp_path / "one"
    second = tmp_path / "two"
    command = [
        sys.executable,
        "-m",
        "backend.app.etl.history.synthetic_publication",
        "generate",
        "--seed",
        "41",
        "--output-dir",
    ]
    for target in (first, second):
        subprocess.run([*command, str(target)], check=True, capture_output=True, text=True)
    files = {p.name: p.read_bytes() for p in first.iterdir()}
    assert files == {p.name: p.read_bytes() for p in second.iterdir()}
    assert len([n for n in files if n.endswith(".csv")]) == 2
    assert len([n for n in files if n.endswith(".xlsx")]) == 2
    for name, data in files.items():
        if name.endswith(".xlsx"):
            season = name.removeprefix("SYNTHETIC_ONLY_").removesuffix(".xlsx").replace("_", "-")
            publication.check_xlsx(data, publication.generate(41), season)
            book = load_workbook(io.BytesIO(data))
            assert tuple(c.value for c in book.active[1]) == EXPECTED_HEADERS
            assert book.properties.creator == "SYNTHETIC_ONLY"
            book.close()
    manifest = json.loads(files["SYNTHETIC_ONLY-manifest.json"])
    assert manifest["data_class"] == "SYNTHETIC_ONLY"
    assert manifest["strict_pit_validated"] is False
    assert manifest["source_equivalence_proven"] is False
    with pytest.raises(FileExistsError):
        publication.export(publication.generate(41), first)


@pytest.mark.parametrize(
    "problem", ["hidden", "comment", "formula", "link", "author", "extra_member"]
)
def test_workbook_metadata_and_active_content_are_rejected(problem):
    data = publication.xlsx_bytes(publication.generate())
    if problem == "extra_member":
        stream = io.BytesIO(data)
        with zipfile.ZipFile(stream, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"SYNTHETIC_ONLY")
        damaged = stream.getvalue()
    else:
        book = load_workbook(io.BytesIO(data))
        if problem == "hidden":
            book.create_sheet("HIDDEN").sheet_state = "hidden"
        if problem == "comment":
            book.active["A2"].comment = Comment("PRIVATE_NOTE", "PRIVATE_AUTHOR")
        if problem == "formula":
            book.active["A2"] = "=1+1"
        if problem == "link":
            book.active["A2"].hyperlink = "https://invalid.example"
        if problem == "author":
            book.properties.creator = "PRIVATE_AUTHOR"
        out = io.BytesIO()
        book.save(out)
        book.close()
        damaged = out.getvalue()
    with pytest.raises(ValueError, match="UNSAFE_PUBLIC_WORKBOOK"):
        publication.check_xlsx(damaged, publication.generate())


def test_clean_workbook_with_relabelled_private_cells_is_rejected():
    package = publication.generate()
    data = publication.xlsx_bytes(package)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                content = content.replace(b"SYNTHETIC_FARM_1", b"PRIVATE_FARM")
            target.writestr(name, content)
    with pytest.raises(ValueError, match="CONTENT_BINDING"):
        publication.check_xlsx(output.getvalue(), package)


def test_csv_identity_and_weights_are_bound_before_export(tmp_path):
    package = publication.generate()
    output = tmp_path / "synthetic"
    publication.export(package, output)
    data = (output / "SYNTHETIC_ONLY_2039_2040.csv").read_bytes()
    publication.check_csv(data, package, "2039-2040")
    with pytest.raises(ValueError, match="CONTENT_BINDING"):
        publication.check_csv(
            data.replace(b"SYNTHETIC_FARM_1", b"PRIVATE_FARM"), package, "2039-2040"
        )


def test_committed_clean_examples_are_exactly_generator_output(tmp_path):
    root = Path(__file__).resolve().parents[3]
    expected = root / "docs/tech-debt/td01c-synthetic-example"
    output = tmp_path / "generated"
    publication.export(publication.generate(41), output)
    assert {f.name: f.read_bytes() for f in output.iterdir()} == {
        f.name: f.read_bytes() for f in expected.iterdir()
    }


def test_path_only_audit_registry_and_original_ref_inventory_are_bound():
    root = Path(__file__).resolve().parents[3] / "docs/tech-debt"
    report = json.loads((root / "td01c-public-derived-audit-summary.json").read_text())
    register = root / "td01c-public-derived-risk-register.csv"
    assert hashlib.sha256(register.read_bytes()).hexdigest() == report["register_sha256"]
    with register.open() as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == report["candidate_paths"]
    assert [r["PATH"] for r in rows] == sorted(set(r["PATH"] for r in rows))
    for row in rows:
        assert set(row) == {
            "PATH",
            "ASSET_CLASS",
            "DISCLOSURE_CATEGORY",
            "HISTORICAL_CONTRACT_DEPENDENCY",
            "CURRENT_PUBLIC_REFERENCE",
            "RECOMMENDED_DISPOSITION",
            "OWNER_DECISION_REQUIRED",
        }
        assert not row["PATH"].startswith("/") and ".." not in Path(row["PATH"]).parts
    delta = json.loads((root / "td01c-ref-inventory-delta.json").read_text())
    assert delta["previous_inventory_rows"] == 742
    assert (
        hashlib.sha256((root / "td01b-public-ref-exposure-r1.csv").read_bytes()).hexdigest()
        == delta["previous_inventory_sha256"]
    )
    assert report["raw_workbooks_read"] is False
