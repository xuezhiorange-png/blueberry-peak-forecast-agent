"""V0_20 source projection stays tied to its approved historical Git blobs.

Current test/production worktree hashes never stand in for historical bytes.
Only public contracts and SYNTHETIC fixtures are read here.
"""

import base64
import hashlib
import io
import json
import subprocess

from openpyxl import load_workbook

from backend.tests.agronomic_history.v0_20_assets import BASE, ROOT, SOURCES, synthetic, vocabulary


def test_v0_20_vocabulary_preserves_original_registry_and_git_identity():
    published = json.loads(
        (ROOT / "frontend/src/dashboard/agronomic/v0_20_vocabulary.json").read_bytes()
    )
    assert published == vocabulary()
    for path in SOURCES:
        raw = subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT)
        assert hashlib.sha256(raw).hexdigest() == published["sources"][path]
    events = json.loads(subprocess.check_output(["git", "show", f"{BASE}:{SOURCES[0]}"], cwd=ROOT))
    assert [r["code"] for r in published["event_types"]] == [
        r["event_type"] for r in events["event_types"]
    ]
    assert published["production_systems"] == [
        "DECIDUOUS_NATURAL",
        "DECIDUOUS_FORCING",
        "EVERGREEN",
    ]


def test_v0_20_synthetic_assets_are_exact_repeatable_two_season_workbooks():
    published = json.loads(
        (ROOT / "frontend/src/dashboard/agronomic/v0_20_synthetic.json").read_bytes()
    )
    for scenario in ("wide", "long", "negative"):
        assert published[scenario] == synthetic(scenario)
        for format_, asset in published[scenario]["assets"].items():
            raw = base64.b64decode(asset["base64"])
            assert hashlib.sha256(raw).hexdigest() == asset["sha256"]
            if format_ == "xlsx":
                wb = load_workbook(io.BytesIO(raw), read_only=True)
                try:
                    assert {"说明", "种植面积", "农事操作", "开花物候", "Excel时间精度"} <= set(
                        wb.sheetnames
                    )
                    cells = [str(c.value) for s in wb for row in s for c in row]
                    assert "2023-2024" in cells and "2024-2025" in cells
                    assert not any(c.data_type == "f" for s in wb for row in s for c in row)
                finally:
                    wb.close()
        assert len(published[scenario]["scope_reference"]) == 3


def test_v0_20_contract_keeps_preparation_separate_from_source_model_and_release_permissions():
    path = ROOT / "docs/v0-20/evidence/v0.20.0-foundation-contract-r1.json"
    contract = json.loads(path.read_bytes())
    assert contract["task_id"] == "V0_20_AGRONOMIC_HISTORY_DATA_FOUNDATION_R1"
    assert contract["base_main_sha"] == BASE
    assert contract["upstream"] == vocabulary()
    assert len(contract["outputs"]) == 5
    assert contract["expected_synthetic"]["total"] == 90
    assert contract["expected_synthetic"]["research_admitted_count"] == 0
    assert all(v is False for v in contract["restrictions"].values())
    for stage in ("s0", "s1", "s2", "s3"):
        assert list((ROOT / "docs/v0-20" / stage).glob("*.md"))
    for scenario, formats in contract["synthetic_assets"].items():
        assets = synthetic(scenario)["assets"]
        for format_, identity in formats.items():
            assert identity["sha256"] == assets[format_]["sha256"]
            assert identity["view_sha256"] == assets[format_]["view_sha256"]
